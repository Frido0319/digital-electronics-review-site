from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path

import fitz
from docx import Document

from . import config


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def relative_asset_path(path: Path) -> str:
    return path.resolve().relative_to(config.PROJECT_ROOT.resolve()).as_posix()


def extract_pdf_text(pdf_path: Path) -> dict:
    doc = fitz.open(str(pdf_path))
    pages = []
    for index in range(doc.page_count):
        text = clean_text(doc[index].get_text("text"))
        pages.append({"page": index + 1, "text": text})
    doc.close()
    return {
        "file": str(pdf_path.relative_to(config.SOURCE_ROOT)),
        "type": "pdf",
        "pages": pages,
        "page_count": len(pages),
    }


def render_pdf_page(pdf_path: Path, page_number: int, output_path: Path, zoom: float = 1.6) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(pdf_path))
    page = doc[page_number - 1]
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    pixmap.save(str(output_path))
    doc.close()


def extract_docx_text_and_media(docx_path: Path, media_dir: Path) -> dict:
    media_dir.mkdir(parents=True, exist_ok=True)
    document = Document(str(docx_path))
    paragraphs = [clean_text(paragraph.text) for paragraph in document.paragraphs if clean_text(paragraph.text)]
    table_rows = []
    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if cells:
                table_rows.append(" | ".join(cells))

    media = []
    with zipfile.ZipFile(docx_path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            source_name = Path(name).name
            target = media_dir / f"{docx_path.stem}_{source_name}"
            with archive.open(name) as source, target.open("wb") as dest:
                shutil.copyfileobj(source, dest)
            media.append(relative_asset_path(target))

    return {
        "file": str(docx_path.relative_to(config.SOURCE_ROOT)),
        "type": "docx",
        "paragraphs": paragraphs,
        "table_rows": table_rows,
        "media": media,
        "text_char_count": sum(len(item) for item in paragraphs + table_rows),
    }


def extract_all_sources() -> dict:
    config.EXTRACTED_TEXT_DIR.mkdir(parents=True, exist_ok=True)
    config.HOMEWORK_IMAGES_DIR.mkdir(parents=True, exist_ok=True)

    course_records = []
    for files in config.SOURCE_FILES.values():
        for file_path in sorted(files):
            if file_path.suffix.lower() == ".pdf":
                course_records.append(extract_pdf_text(file_path))
            else:
                course_records.append(
                    {
                        "file": str(file_path.relative_to(config.SOURCE_ROOT)),
                        "type": file_path.suffix.lower().lstrip("."),
                        "note": "旧版 PPT 需在后续任务中转换或截图。",
                    }
                )

    homework_records = [
        extract_docx_text_and_media(path, config.HOMEWORK_IMAGES_DIR)
        for path in config.HOMEWORK_FILES
    ]

    result = {"course": course_records, "homework": homework_records}
    (config.EXTRACTED_TEXT_DIR / "sources.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return result
